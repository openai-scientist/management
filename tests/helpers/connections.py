import asyncio
import io
from collections.abc import AsyncIterator, Awaitable, Callable
from contextlib import asynccontextmanager
from datetime import UTC, date, datetime, timedelta
from urllib.parse import parse_qs, urlsplit
from uuid import uuid4

import asyncpg
import openpyxl
from cryptography import x509
from cryptography.hazmat.primitives import hashes, serialization
from cryptography.hazmat.primitives.asymmetric import ec
from cryptography.x509.oid import NameOID
from httpx import AsyncClient
from openpyxl.chart import BarChart, Reference
from sqlalchemy import select

from management.connections.connectors import build_connector_factory
from management.connections.connectors.base import Column
from management.connections.google_drive_oauth import GoogleGrant
from management.connections.models import DataConnection, GoogleConnectionGrant
from management.infra.secret_box import SecretBox
from tests.fakes import FakeGoogleSheets, FakeInfluxDB, FakePrometheus, FakeSpreadsheet, FakeTable
from tests.helpers.harness import CONNECTION_KEY, Harness
from tests.helpers.identity import login, mutation_headers
from tests.helpers.projects import PROJECTS, create_project

SECRET = "s3cret-pass-w0rd"


def connection_body(name: str = "Warehouse", **config: object) -> dict:
    return {
        "name": name,
        "kind": "postgres",
        "config": {
            "host": "db.example.com",
            "database": "analytics",
            "username": "reader",
            **config,
        },
        "secret": {"password": SECRET},
    }


async def create_connection(client: AsyncClient, session: dict, project_id: str, **kwargs):
    return await client.post(
        f"{PROJECTS}/{project_id}/connections",
        json=connection_body(**kwargs),
        headers=mutation_headers(session["csrf_token"]),
    )


async def wait_until(condition) -> None:
    for _ in range(200):
        if condition():
            return
        await asyncio.sleep(0.005)
    raise AssertionError("condition was not reached")


PROMETHEUS_URL = "https://metrics.example.com:9090/prometheus"


# A public address: the guard refuses the ranges kept for examples. Nothing is sent to it.
IP = "93.184.216.34"
PROMETHEUS_SERIES = {
    "type": "timeseries",
    "name": "latency",
    "fields": [],
    "tags": ["host"],
    "start": "2026-10-01T00:00:00Z",
    "end": "2026-10-01T03:00:00Z",
    "bucket": "1h",
    "aggregate": "mean",
}


def with_prometheus(harness: Harness) -> FakePrometheus:
    """Put a fake Prometheus behind the real connector factory; one metric is in it."""
    server = FakePrometheus(base_path="/prometheus")
    server.add("latency", {"host": "a"}, [(at(0.25), 1.0), (at(0.5), 2.0), (at(1.5), 10.0)])
    server.add("latency", {"host": "b"}, [(at(0.75), 5.0), (at(2.5), 2.5)])
    server.add("http_requests_total", {"code": "200"}, [])

    async def resolver(host: str, port: int) -> list[str]:
        return {"metrics.example.com": [IP], "internal.example.com": ["10.0.0.5"]}[host]

    harness.app.state.connector_factory = build_connector_factory(
        harness.settings, resolver, http_transport=server.transport
    )
    return server


def prometheus_body(name: str = "Metrics", url: str = PROMETHEUS_URL, **secret: str) -> dict:
    return {"name": name, "kind": "prometheus", "config": {"url": url}, "secret": secret}


async def create(client: AsyncClient, session: dict, project_id: str, body: dict):
    return await client.post(
        f"{PROJECTS}/{project_id}/connections",
        json=body,
        headers=mutation_headers(session["csrf_token"]),
    )


async def connected_prometheus(harness: Harness, client: AsyncClient, **secret: str):
    """A signed-in owner, their project and the URL of a saved Prometheus connection."""
    session = await login(harness, client, uid="owner", email="owner@example.com")
    project = await create_project(client, session)
    created = await create(client, session, project["id"], prometheus_body(**secret))
    assert created.status_code == 201, created.text
    url = f"{PROJECTS}/{project['id']}/connections/{created.json()['data']['id']}"
    return session, project, url


async def connections(harness: Harness) -> list[DataConnection]:
    async with harness.factory() as db:
        return list(await db.scalars(select(DataConnection)))


ORDERS = FakeTable(
    columns=[Column("id", "integer"), Column("note", "text"), Column("paid", "boolean")],
    rows=[(n, f"note {n}", n % 2 == 0) for n in range(1, 151)],
)


SERIES = {
    "type": "timeseries",
    "name": "orders",
    "fields": ["id"],
    "tags": [],
    "start": "2026-10-01T00:00:00Z",
    "end": "2026-10-02T00:00:00Z",
    "bucket": "1h",
    "aggregate": "mean",
}


def external_database(harness: Harness) -> None:
    harness.connectors.tables = {
        ("public", "orders"): ORDERS,
        ("public", "order_lines"): FakeTable(columns=[Column("id", "integer")]),
        ("public", "customers"): FakeTable(
            columns=[Column("id", "integer"), Column("bio", "text")],
            rows=[(1, "x" * 600), (2, None)],
        ),
        ("sales", "regions"): FakeTable(columns=[Column("code", "text")], rows=[("north",)]),
    }


async def connected_project(harness: Harness, client: AsyncClient, **who: str):
    """A signed-in owner, their project and the URL of a saved connection in it."""
    session = await login(
        harness, client, **(who or {"uid": "owner", "email": "owner@example.com"})
    )
    project = await create_project(client, session)
    created = await create_connection(client, session, project["id"])
    assert created.status_code == 201, created.text
    url = f"{PROJECTS}/{project['id']}/connections/{created.json()['data']['id']}"
    return session, project, url


async def preview(client: AsyncClient, session: dict, url: str, source: dict):
    return await client.post(
        f"{url}/preview", json={"source": source}, headers=mutation_headers(session["csrf_token"])
    )


def failure(response) -> tuple[int, str, str | None]:
    error = response.json()["error"]
    return response.status_code, error["code"], error.get("reason")


STATE_COOKIE = "platform_google_drive_state"
CALLBACK = "/api/v1/connections/google/callback"


def start_path(project_id: str) -> str:
    return f"{PROJECTS}/{project_id}/connections/google/start"


async def start(client, project_id: str) -> str:
    """Leave for Google from this browser; the state Google will hand back."""
    started = await client.get(start_path(project_id))
    assert started.status_code == 302, started.text
    return parse_qs(urlsplit(started.headers["location"]).query)["state"][0]


def workbook() -> bytes:
    """A workbook as Excel saves one: two sheets of cells and a sheet that is a chart."""
    book = openpyxl.Workbook()
    scores = book.active
    scores.title = "Scores"
    scores.append(["student", "score", "taken", "day", "passed", "double", "group"])
    scores.append(["An", 7.5, datetime(2026, 3, 1, 9, 30), date(2026, 3, 1), True, "=B2*2", "x"])
    scores.append(["Binh", 8, None, None, False, None, None])
    # Only the first cell of a merged range holds its value.
    scores.merge_cells("G2:G3")
    book.create_sheet("Second").append(["only"])
    book["Second"].append([1])
    chart = BarChart()
    chart.add_data(Reference(scores, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    book.create_chartsheet("Chart").add_chart(chart)
    saved = io.BytesIO()
    book.save(saved)
    return saved.getvalue()


SPREADSHEET_ID = "1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms"
SPREADSHEET_URL = f"https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/edit?gid=0#gid=0"
PRIVATE_ID = "1-a-spreadsheet-of-somebody-else-entirely"
TITLE = "Survey 2026"
SUBJECT = "google-sub-1"
EMAIL = "dat@gmail.com"
REFRESH_TOKEN = "refresh-token-that-must-never-be-shown"


def with_sheets(harness: Harness) -> FakeGoogleSheets:
    """Put a fake Sheets API behind the real connector factory; one spreadsheet is in it."""
    sheets = FakeGoogleSheets()
    sheets.spreadsheets[SPREADSHEET_ID] = FakeSpreadsheet(
        title=TITLE,
        tabs={
            "Answers": [
                ["student", "school", "score"],
                ["An", "A", 7.5],
                ["Binh", "B, north", 8],
                ["Chi", "", 6.25],
            ],
            "Notes": [["note"], ["first"]],
            "Broken": [["a", ""], [1, "stray"]],
        },
        readers={REFRESH_TOKEN},
    )
    sheets.spreadsheets[PRIVATE_ID] = FakeSpreadsheet(title="Private", tabs={}, readers=set())
    harness.app.state.connector_factory = build_connector_factory(
        harness.settings, google_oauth=harness.google_drive, google_transport=sheets.transport
    )
    return sheets


async def google_grant(
    harness: Harness,
    client: AsyncClient,
    project_id: str,
    *,
    refresh_token: str = REFRESH_TOKEN,
    subject: str = SUBJECT,
) -> str:
    """Give access as a Google account from this browser; the id the frontend is handed."""
    harness.google_drive.grant = GoogleGrant(refresh_token, subject, EMAIL)
    state = await start(client, project_id)
    done = await client.get(CALLBACK, params={"code": "one-use-code", "state": state})
    return parse_qs(urlsplit(done.headers["location"]).query)["google_grant"][0]


def sheets_body(grant_id: str, name: str = "Survey", spreadsheet: str = SPREADSHEET_URL) -> dict:
    return {
        "name": name,
        "kind": "google_sheets",
        "config": {"spreadsheet": spreadsheet},
        "grant_id": grant_id,
    }


async def create_sheets_connection(
    client: AsyncClient, session: dict, project_id: str, grant_id: str, **body: str
):
    return await client.post(
        f"{PROJECTS}/{project_id}/connections",
        json=sheets_body(grant_id, **body),
        headers=mutation_headers(session["csrf_token"]),
    )


async def reauthorize(client: AsyncClient, session: dict, url: str, grant_id: str):
    return await client.post(
        f"{url}/reauthorize",
        json={"grant_id": grant_id},
        headers=mutation_headers(session["csrf_token"]),
    )


async def grant_ids(harness: Harness) -> set[str]:
    async with harness.factory() as db:
        return {str(grant_id) for grant_id in await db.scalars(select(GoogleConnectionGrant.id))}


def stored_token(row: DataConnection) -> str:
    return SecretBox(CONNECTION_KEY).open(row.secret_ciphertext)["refresh_token"]


async def one_grant_makes_one_connection(harness: Harness) -> None:
    """Two requests with one grant, both past the first check before either saves.

    For a harness on PostgreSQL: the in-memory SQLite one runs every request on a single
    database connection, where one request's rollback undoes the other's work.
    """
    sheets = with_sheets(harness)
    sheets.hold = asyncio.Event()
    async with harness.client() as client:
        session = await login(harness, client, uid="owner", email="owner@example.com")
        project = await create_project(client, session)
        grant_id = await google_grant(harness, client, project["id"])

        # Both find the grant unused, and both are still asking Google about the spreadsheet.
        attempts = [
            asyncio.create_task(
                create_sheets_connection(client, session, project["id"], grant_id, name=name)
            )
            for name in ("First", "Second")
        ]
        await wait_until(lambda: len(sheets.requests) == 2)
        sheets.hold.set()
        answers = await asyncio.gather(*attempts)

    assert sorted(answer.status_code for answer in answers) == [201, 422]
    refused = next(answer for answer in answers if answer.status_code == 422)
    assert failure(refused) == (422, "GOOGLE_GRANT_INVALID", None)
    assert len(await connections(harness)) == 1
    assert await grant_ids(harness) == set()


INFLUXDB_URL = "https://influx.example.com:8086/influx"


def with_influxdb(harness: Harness) -> FakeInfluxDB:
    """Put a fake InfluxDB behind the real connector factory; the points of the fake
    Prometheus are in it, as the field `value` of one measurement."""
    server = FakeInfluxDB(base_path="/influx")
    for host, moment, value in [
        ("a", 0.25, 1.0),
        ("a", 0.5, 2.0),
        ("a", 1.5, 10.0),
        ("b", 0.75, 5.0),
        ("b", 2.5, 2.5),
    ]:
        server.write("latency", {"host": host}, {"value": value, "note": "ok"}, at(moment))
    server.write("jobs", {}, {"done": 3}, at(0.5))

    async def resolver(host: str, port: int) -> list[str]:
        return {"influx.example.com": [IP], "internal.example.com": ["10.0.0.5"]}[host]

    harness.app.state.connector_factory = build_connector_factory(
        harness.settings, resolver, http_transport=server.transport
    )
    return server


def influxdb_body(
    name: str = "Benchmarks", url: str = INFLUXDB_URL, database: str = "bench", **secret: str
) -> dict:
    return {
        "name": name,
        "kind": "influxdb",
        "config": {"url": url, "database": database},
        "secret": secret,
    }


async def connected_influxdb(harness: Harness, client: AsyncClient, **secret: str):
    """A signed-in owner, their project and the URL of a saved InfluxDB connection."""
    session = await login(harness, client, uid="owner", email="owner@example.com")
    project = await create_project(client, session)
    created = await create(client, session, project["id"], influxdb_body(**secret))
    assert created.status_code == 201, created.text
    url = f"{PROJECTS}/{project['id']}/connections/{created.json()['data']['id']}"
    return session, project, url


DAY = datetime(2026, 10, 1, tzinfo=UTC)


def at(hours: float = 0, **more: float) -> int:
    """A moment of the day the tests read, in milliseconds as a sample carries it."""
    return round((DAY + timedelta(hours=hours, **more)).timestamp() * 1000)


Handler = Callable[[asyncio.StreamReader, asyncio.StreamWriter], Awaitable[None]]


@asynccontextmanager
async def tcp_server(handler: Handler) -> AsyncIterator[int]:
    server = await asyncio.start_server(handler, "127.0.0.1", 0)
    try:
        yield server.sockets[0].getsockname()[1]
    finally:
        server.close()


def self_signed_certificate(directory, dns_name: str) -> tuple[str, str]:
    key = ec.generate_private_key(ec.SECP256R1())
    name = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, dns_name)])
    now = datetime.now(UTC)
    certificate = (
        x509.CertificateBuilder()
        .subject_name(name)
        .issuer_name(name)
        .public_key(key.public_key())
        .serial_number(x509.random_serial_number())
        .not_valid_before(now - timedelta(days=1))
        .not_valid_after(now + timedelta(days=1))
        .add_extension(x509.SubjectAlternativeName([x509.DNSName(dns_name)]), critical=False)
        .add_extension(x509.BasicConstraints(ca=True, path_length=None), critical=True)
        .sign(key, hashes.SHA256())
    )
    cert_path, key_path = directory / "server.pem", directory / "server.key"
    cert_path.write_bytes(certificate.public_bytes(serialization.Encoding.PEM))
    key_path.write_bytes(
        key.private_bytes(
            serialization.Encoding.PEM,
            serialization.PrivateFormat.PKCS8,
            serialization.NoEncryption(),
        )
    )
    return str(cert_path), str(key_path)


@asynccontextmanager
async def direct(url) -> AsyncIterator[asyncpg.Connection]:
    """A connection of the test's own, to set a scene and to look at what is left of it."""
    connection = await asyncpg.connect(
        host=url.host,
        port=url.port or 5432,
        user=url.username,
        password=url.password,
        database=url.database,
    )
    try:
        yield connection
    finally:
        await connection.close()


@asynccontextmanager
async def sample_schema(url) -> AsyncIterator[str]:
    schema = f"conn_test_{uuid4().hex}"
    async with direct(url) as admin:
        await admin.execute(
            f"""
            CREATE SCHEMA {schema};
            SET search_path = {schema};
            CREATE TABLE orders (
                id integer PRIMARY KEY, note text, amount numeric(10, 2), tags text[],
                payload jsonb, placed_on date
            );
            INSERT INTO orders
            SELECT n, 'note ' || n, n * 1.5, ARRAY['a', 'b'], '{{"k": 1}}', DATE '2026-01-01' + n
            FROM generate_series(1, 250) AS n;
            CREATE VIEW big_orders AS SELECT id, amount FROM orders WHERE amount > 100;
            CREATE MATERIALIZED VIEW order_totals AS SELECT count(*) AS orders FROM orders;
            CREATE TABLE order_100_percent (id integer);
            """
        )
        await admin.execute('CREATE TABLE "we""ird name" ("the ""col""" text)')
        await admin.execute("""INSERT INTO "we""ird name" VALUES ('found')""")
        try:
            yield schema
        finally:
            await admin.execute(f"DROP SCHEMA {schema} CASCADE")
